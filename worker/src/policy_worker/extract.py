"""Text extraction — same interface as the harvester's extract module.

TODO: replace with a dependency on the published policy-harvester package once
both repos are on a registry; until then this is a deliberate, small copy kept
in sync by the shared test corpus.
"""

from __future__ import annotations

from importlib.metadata import version
from pathlib import Path


def extract_text(path: str | Path) -> tuple[str | None, str | None]:
    """Extract plain text; returns (text, tool) or (None, None) if unsupported."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _pdf(path)
    if suffix == ".docx":
        return _docx(path)
    if suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="replace"), "plaintext"
    return None, None


def _pdf(path: Path) -> tuple[str | None, str]:
    import pymupdf

    tool = f"pymupdf/{version('pymupdf')}"
    with pymupdf.open(path) as doc:
        text = "\n".join(page.get_text() for page in doc)
    return text or None, tool


def _docx(path: Path) -> tuple[str | None, str]:
    import docx

    tool = f"python-docx/{version('python-docx')}"
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    text = "\n".join(part for part in parts if part.strip())
    return text or None, tool
